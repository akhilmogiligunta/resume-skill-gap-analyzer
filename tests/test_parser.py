from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from src.resume_parser import (
    get_file_extension,
    is_supported_file,
    extract_text,
    extract_text_from_pdf,
    extract_text_from_docx,
)


def test_pdf_extension():

    assert get_file_extension(
        "resume.pdf"
    ) == ".pdf"


def test_docx_extension():

    assert get_file_extension(
        "resume.docx"
    ) == ".docx"


def test_unsupported_extension():

    assert not is_supported_file(
        "resume.exe"
    )


def test_supported_files():

    assert is_supported_file(
        "resume.pdf"
    )

    assert is_supported_file(
        "resume.docx"
    )


def test_pdf_text_extraction(tmp_path):

    pdf_path = tmp_path / "resume.pdf"

    pdf = canvas.Canvas(
        str(pdf_path)
    )

    pdf.drawString(
        100,
        750,
        "Python SQL Machine Learning"
    )

    pdf.save()

    text = extract_text_from_pdf(
        pdf_path
    )

    assert "Python" in text
    assert "SQL" in text
    assert "Machine Learning" in text


def test_docx_text_extraction(tmp_path):

    docx_path = tmp_path / "resume.docx"

    document = Document()

    document.add_paragraph(
        "Python React Docker"
    )

    document.save(
        docx_path
    )

    text = extract_text_from_docx(
        docx_path
    )

    assert "Python React Docker" in text


def test_extract_text_pdf(tmp_path):

    pdf_path = tmp_path / "resume.pdf"

    pdf = canvas.Canvas(
        str(pdf_path)
    )

    pdf.drawString(
        100,
        750,
        "Python Developer"
    )

    pdf.save()

    text = extract_text(
        pdf_path
    )

    assert "Python Developer" in text


def test_extract_text_docx(tmp_path):

    docx_path = tmp_path / "resume.docx"

    document = Document()

    document.add_paragraph(
        "Java Developer"
    )

    document.save(
        docx_path
    )

    text = extract_text(
        docx_path
    )

    assert "Java Developer" in text


def test_extract_text_unsupported_file(tmp_path):

    file_path = tmp_path / "resume.txt"

    file_path.write_text(
        "Python Developer"
    )

    try:

        extract_text(file_path)

        assert False

    except ValueError as error:

        assert "Unsupported file type" in str(error)